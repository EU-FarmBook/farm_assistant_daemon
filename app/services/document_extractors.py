"""
Per-format text extraction for chat file attachments.

One entry point — `extract_text(path, filename)` — dispatches on the file
extension and returns plain text that the shared attachment pipeline
(pdf_service: summarise -> chunk -> per-question relevance injection) consumes
unchanged. Heavy format libraries are lazy-imported so a missing optional
dependency degrades into a clear per-file processing error instead of breaking
service startup.

Legacy binary formats (.doc, .ppt) are deliberately unsupported: there is no
maintained lightweight Python reader for them. The upload endpoint rejects
them with a "save as DOCX/PPTX" message before this module is ever reached.
"""
from __future__ import annotations

import csv
import io
import json
import logging
import re
from pathlib import Path

logger = logging.getLogger("farm-assistant.extractors")

# Extraction caps shared across formats (mirrors the original PDF limits).
MAX_CHARS = 240_000
MAX_PDF_PAGES = 80
MAX_TABLE_ROWS = 500          # per CSV file / spreadsheet sheet
MAX_SHEETS = 10
MAX_SLIDES = 120

# Extensions the /files/document endpoint accepts, with their canonical mime.
SUPPORTED_DOCUMENT_TYPES: dict[str, str] = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".txt": "text/plain",
    ".csv": "text/csv",
    ".json": "application/json",
    ".xls": "application/vnd.ms-excel",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
}

# Rejected with an actionable message instead of a generic "unsupported".
LEGACY_UNSUPPORTED: dict[str, str] = {
    ".doc": "Legacy .doc files are not supported. Please save the file as .docx and upload again.",
    ".ppt": "Legacy .ppt files are not supported. Please save the file as .pptx and upload again.",
}


def _decode_bytes(payload: bytes) -> str:
    for encoding in ("utf-8", "utf-8-sig", "latin-1"):
        try:
            return payload.decode(encoding)
        except UnicodeDecodeError:
            continue
    return payload.decode("utf-8", errors="replace")


def _extract_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader  # type: ignore
    except Exception as e:
        raise RuntimeError("pypdf is required for PDF extraction. Install dependency first.") from e

    reader = PdfReader(str(path))
    pages = reader.pages[:MAX_PDF_PAGES]
    parts: list[str] = []
    total = 0
    for p in pages:
        try:
            txt = (p.extract_text() or "").strip()
        except Exception:
            txt = ""
        if not txt:
            continue
        remain = MAX_CHARS - total
        if remain <= 0:
            break
        txt = txt[:remain]
        parts.append(txt)
        total += len(txt)
    return "\n\n".join(parts).strip()


def _extract_txt(path: Path) -> str:
    return _decode_bytes(path.read_bytes())[:MAX_CHARS].strip()


def _extract_json(path: Path) -> str:
    raw = _decode_bytes(path.read_bytes())
    try:
        parsed = json.loads(raw)
        pretty = json.dumps(parsed, indent=2, ensure_ascii=False)
    except Exception:
        pretty = raw
    return pretty[:MAX_CHARS].strip()


def _rows_to_text(rows: list[list[str]], truncated: bool, label: str) -> str:
    lines = [" | ".join(cell for cell in row) for row in rows]
    if truncated:
        lines.append(f"... ({label} truncated at {MAX_TABLE_ROWS} rows)")
    return "\n".join(lines)


def _extract_csv(path: Path) -> str:
    raw = _decode_bytes(path.read_bytes())
    try:
        dialect = csv.Sniffer().sniff(raw[:4096], delimiters=",;\t|")
    except Exception:
        dialect = csv.excel
    rows: list[list[str]] = []
    truncated = False
    for i, row in enumerate(csv.reader(io.StringIO(raw), dialect)):
        if i >= MAX_TABLE_ROWS:
            truncated = True
            break
        rows.append([str(c).strip() for c in row])
    return _rows_to_text(rows, truncated, "file")[:MAX_CHARS].strip()


def _extract_docx(path: Path) -> str:
    try:
        import docx  # type: ignore  # python-docx
    except Exception as e:
        raise RuntimeError("python-docx is required for DOCX extraction. Install dependency first.") from e

    document = docx.Document(str(path))
    parts: list[str] = []
    total = 0
    for para in document.paragraphs:
        txt = (para.text or "").strip()
        if not txt:
            continue
        parts.append(txt)
        total += len(txt)
        if total >= MAX_CHARS:
            break
    for table in document.tables:
        if total >= MAX_CHARS:
            break
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            line = " | ".join(cells).strip(" |")
            if line:
                parts.append(line)
                total += len(line)
            if total >= MAX_CHARS:
                break
    return "\n".join(parts)[:MAX_CHARS].strip()


def _extract_xlsx(path: Path) -> str:
    try:
        from openpyxl import load_workbook  # type: ignore
    except Exception as e:
        raise RuntimeError("openpyxl is required for XLSX extraction. Install dependency first.") from e

    workbook = load_workbook(str(path), read_only=True, data_only=True)
    parts: list[str] = []
    total = 0
    for sheet in workbook.worksheets[:MAX_SHEETS]:
        rows: list[list[str]] = []
        truncated = False
        for i, row in enumerate(sheet.iter_rows(values_only=True)):
            if i >= MAX_TABLE_ROWS:
                truncated = True
                break
            rows.append(["" if c is None else str(c).strip() for c in row])
        if not any(any(cell for cell in row) for row in rows):
            continue
        block = f"Sheet: {sheet.title}\n{_rows_to_text(rows, truncated, 'sheet')}"
        parts.append(block)
        total += len(block)
        if total >= MAX_CHARS:
            break
    workbook.close()
    return "\n\n".join(parts)[:MAX_CHARS].strip()


def _extract_xls(path: Path) -> str:
    try:
        import xlrd  # type: ignore
    except Exception as e:
        raise RuntimeError("xlrd is required for XLS extraction. Install dependency first.") from e

    workbook = xlrd.open_workbook(str(path))
    parts: list[str] = []
    total = 0
    for sheet in workbook.sheets()[:MAX_SHEETS]:
        rows: list[list[str]] = []
        truncated = sheet.nrows > MAX_TABLE_ROWS
        for i in range(min(sheet.nrows, MAX_TABLE_ROWS)):
            rows.append([str(c).strip() for c in sheet.row_values(i)])
        if not any(any(cell for cell in row) for row in rows):
            continue
        block = f"Sheet: {sheet.name}\n{_rows_to_text(rows, truncated, 'sheet')}"
        parts.append(block)
        total += len(block)
        if total >= MAX_CHARS:
            break
    return "\n\n".join(parts)[:MAX_CHARS].strip()


def _extract_pptx(path: Path) -> str:
    try:
        from pptx import Presentation  # type: ignore  # python-pptx
    except Exception as e:
        raise RuntimeError("python-pptx is required for PPTX extraction. Install dependency first.") from e

    presentation = Presentation(str(path))
    parts: list[str] = []
    total = 0
    for index, slide in enumerate(presentation.slides, start=1):
        if index > MAX_SLIDES:
            break
        texts: list[str] = []
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            txt = (shape.text_frame.text or "").strip()
            if txt:
                texts.append(txt)
        if not texts:
            continue
        block = f"Slide {index}:\n" + "\n".join(texts)
        parts.append(block)
        total += len(block)
        if total >= MAX_CHARS:
            break
    return "\n\n".join(parts)[:MAX_CHARS].strip()


_EXTRACTORS = {
    ".pdf": _extract_pdf,
    ".txt": _extract_txt,
    ".json": _extract_json,
    ".csv": _extract_csv,
    ".docx": _extract_docx,
    ".xlsx": _extract_xlsx,
    ".xls": _extract_xls,
    ".pptx": _extract_pptx,
}


def extract_text(path: Path, filename: str) -> str:
    """Extract plain text from a supported attachment; raises on unsupported/broken files."""
    suffix = Path(filename or "").suffix.lower() or path.suffix.lower()
    extractor = _EXTRACTORS.get(suffix)
    if extractor is None:
        raise RuntimeError(f"Unsupported file type: {suffix or 'unknown'}")
    text = extractor(path)
    # Collapse pathological whitespace early so chunking stays meaningful.
    return re.sub(r"\n{3,}", "\n\n", text or "").strip()
